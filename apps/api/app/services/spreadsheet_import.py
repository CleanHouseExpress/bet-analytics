import hashlib
import json
import re
import unicodedata
from dataclasses import dataclass
from datetime import UTC, date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session


@dataclass(slots=True)
class SpreadsheetImportResult:
    run_id: int
    rows_seen: int = 0
    rows_imported: int = 0
    rows_skipped: int = 0
    rows_failed: int = 0


class SpreadsheetImporter:
    """Import legacy project spreadsheets into the canonical PostgreSQL model."""

    HEADER_ALIASES = {
        "date": {"data", "date", "placed_at", "data_aposta"},
        "home": {"mandante", "home", "home_team", "time_casa"},
        "away": {"visitante", "away", "away_team", "time_fora"},
        "score": {"resultado", "placar", "score", "pontuacao"},
        "round": {"rodada", "round"},
        "competition": {"competicao", "competição", "liga", "competition", "league"},
        "season": {"temporada", "season", "ano"},
        "goals": {"gols_minuto", "minutos_gols", "gols", "goals_timeline"},
        "stake": {"stake", "valor_apostado", "valor_investido", "aposta", "valor"},
        "odd": {"odd", "odds", "cotacao", "cotação", "total_odd"},
        "return": {"retorno", "ganhos", "return", "return_amount"},
        "result": {"resultado_aposta", "status_aposta", "result", "resultado"},
        "market": {"mercado", "market", "tipo_mercado"},
        "selection": {"selecao", "seleção", "selection", "palpite", "posicao", "posição"},
        "bet_id": {"id_aposta", "bet_id", "id"},
        "is_live": {"live", "ao_vivo", "is_live"},
        "minute": {"minuto", "minute", "minuto_aposta"},
        "movement_type": {"tipo_movimento", "movimento", "movement_type", "tipo"},
        "balance": {"saldo", "saldo_global", "balance", "banca"},
    }

    def __init__(self, db: Session, account_name: str = "Betano") -> None:
        self.db = db
        self.account_name = account_name

    def import_file(
        self,
        path: str | Path,
        *,
        default_competition: str = "Imported Historical Data",
        default_season: str = "legacy",
    ) -> SpreadsheetImportResult:
        file_path = Path(path)
        file_hash = hashlib.sha256(file_path.read_bytes()).hexdigest()
        run_id = self._start_run(file_path.name, file_hash)
        result = SpreadsheetImportResult(run_id=run_id)
        workbook = load_workbook(file_path, data_only=True, read_only=True)

        try:
            account_id = self._ensure_account()
            for worksheet in workbook.worksheets:
                rows = worksheet.iter_rows(values_only=True)
                header_row = next(rows, None)
                if not header_row:
                    continue
                headers = [self._normalize(value) for value in header_row]
                canonical = self._canonical_header_map(headers)
                sheet_kind = self._classify_sheet(canonical)

                for row_number, values in enumerate(rows, start=2):
                    if not any(value is not None and str(value).strip() for value in values):
                        continue
                    result.rows_seen += 1
                    payload = self._row_payload(headers, values)
                    fingerprint = self._fingerprint(worksheet.title, payload)
                    if self._fingerprint_exists(fingerprint):
                        result.rows_skipped += 1
                        continue
                    try:
                        entity_type, entity_id, imported = self._import_row(
                            sheet_kind,
                            canonical,
                            payload,
                            account_id,
                            default_competition,
                            default_season,
                        )
                        status = "imported" if imported else "skipped"
                        self._record_row(
                            run_id,
                            worksheet.title,
                            row_number,
                            fingerprint,
                            status,
                            payload,
                            entity_type,
                            entity_id,
                        )
                        if imported:
                            result.rows_imported += 1
                        else:
                            result.rows_skipped += 1
                    except Exception as exc:
                        self._record_row(
                            run_id,
                            worksheet.title,
                            row_number,
                            fingerprint,
                            "failed",
                            payload,
                            error=str(exc),
                        )
                        result.rows_failed += 1
            self._finish_run(result, "success")
            self.db.commit()
            return result
        except Exception:
            self.db.rollback()
            self._finish_run(result, "failed")
            self.db.commit()
            raise
        finally:
            workbook.close()

    def _import_row(
        self,
        sheet_kind: str,
        canonical: dict[str, str],
        payload: dict[str, Any],
        account_id: int,
        default_competition: str,
        default_season: str,
    ) -> tuple[str | None, int | None, bool]:
        if sheet_kind == "match":
            match_id = self._import_match(canonical, payload, default_competition, default_season)
            return "match", match_id, True
        if sheet_kind == "bet":
            bet_id = self._import_bet(canonical, payload, account_id)
            return ("bet", bet_id, bet_id is not None)
        if sheet_kind == "movement":
            movement_id = self._import_movement(canonical, payload, account_id)
            return ("bankroll_movement", movement_id, movement_id is not None)
        return None, None, False

    def _import_match(
        self,
        canonical: dict[str, str],
        row: dict[str, Any],
        default_competition: str,
        default_season: str,
    ) -> int:
        home = self._value(row, canonical, "home")
        away = self._value(row, canonical, "away")
        if not home or not away:
            raise ValueError("historical match row requires home and away teams")
        competition_name = self._value(row, canonical, "competition") or default_competition
        season_name = str(self._value(row, canonical, "season") or default_season)
        kickoff = self._as_datetime(self._value(row, canonical, "date"))
        if kickoff is None:
            raise ValueError("historical match row requires a date")
        competition_id = self._ensure_competition(str(competition_name))
        season_id = self._ensure_season(competition_id, season_name)
        home_id = self._ensure_team(str(home))
        away_id = self._ensure_team(str(away))
        home_score, away_score = self._parse_score(self._value(row, canonical, "score"))
        match_id = self.db.execute(
            text(
                """
                INSERT INTO matches (
                    competition_id, season_id, home_team_id, away_team_id,
                    kickoff_at, status, home_score, away_score, created_at, updated_at
                ) VALUES (
                    :competition_id, :season_id, :home_team_id, :away_team_id,
                    :kickoff_at, :status, :home_score, :away_score, now(), now()
                )
                ON CONFLICT (season_id, home_team_id, away_team_id, kickoff_at)
                DO UPDATE SET home_score = EXCLUDED.home_score,
                              away_score = EXCLUDED.away_score,
                              status = EXCLUDED.status,
                              updated_at = now()
                RETURNING id
                """
            ),
            {
                "competition_id": competition_id,
                "season_id": season_id,
                "home_team_id": home_id,
                "away_team_id": away_id,
                "kickoff_at": kickoff,
                "status": "finished" if home_score is not None else "scheduled",
                "home_score": home_score,
                "away_score": away_score,
            },
        ).scalar_one()
        goals = self._value(row, canonical, "goals")
        if goals:
            self._import_goal_timeline(match_id, str(goals), home_id, away_id, str(home), str(away))
        return match_id

    def _import_bet(
        self, canonical: dict[str, str], row: dict[str, Any], account_id: int
    ) -> int | None:
        stake = self._as_decimal(self._value(row, canonical, "stake"))
        if stake is None:
            return None
        raw_result = self._value(row, canonical, "result")
        result = self._normalize_bet_result(raw_result)
        return_amount = self._as_decimal(self._value(row, canonical, "return"))
        if return_amount is not None and return_amount == stake and result is None:
            return None
        placed_at = self._as_datetime(self._value(row, canonical, "date")) or datetime.now(UTC)
        odd = self._as_decimal(self._value(row, canonical, "odd"))
        external_id = self._value(row, canonical, "bet_id")
        market = str(self._value(row, canonical, "market") or "legacy")
        selection = str(self._value(row, canonical, "selection") or market)
        profit_loss = None if return_amount is None else return_amount - stake
        fingerprint = hashlib.sha256(
            f"bet|{external_id}|{placed_at.isoformat()}|{stake}|{odd}|{selection}".encode()
        ).hexdigest()
        bet_id = self.db.execute(
            text(
                """
                INSERT INTO bets (
                    account_id, external_bet_id, placed_at, bet_type, stake, total_odd,
                    status, result, return_amount, profit_loss, source, source_fingerprint,
                    created_at, updated_at
                ) VALUES (
                    :account_id, :external_bet_id, :placed_at, 'single', :stake, :odd,
                    'settled', :result, :return_amount, :profit_loss, 'spreadsheet',
                    :fingerprint, now(), now()
                )
                ON CONFLICT (source_fingerprint)
                DO UPDATE SET result = EXCLUDED.result,
                              return_amount = EXCLUDED.return_amount,
                              profit_loss = EXCLUDED.profit_loss,
                              updated_at = now()
                RETURNING id
                """
            ),
            {
                "account_id": account_id,
                "external_bet_id": str(external_id) if external_id else None,
                "placed_at": placed_at,
                "stake": stake,
                "odd": odd,
                "result": result,
                "return_amount": return_amount,
                "profit_loss": profit_loss,
                "fingerprint": fingerprint,
            },
        ).scalar_one()
        self.db.execute(text("DELETE FROM bet_selections WHERE bet_id = :bet_id"), {"bet_id": bet_id})
        self.db.execute(
            text(
                """
                INSERT INTO bet_selections (
                    bet_id, market, selection, odd, is_live, minute_placed, metadata
                ) VALUES (
                    :bet_id, :market, :selection, :odd, :is_live, :minute_placed,
                    CAST(:metadata AS jsonb)
                )
                """
            ),
            {
                "bet_id": bet_id,
                "market": market,
                "selection": selection,
                "odd": odd,
                "is_live": self._as_bool(self._value(row, canonical, "is_live")),
                "minute_placed": self._as_int(self._value(row, canonical, "minute")),
                "metadata": json.dumps(row, default=str),
            },
        )
        if result and return_amount is not None:
            self.db.execute(
                text(
                    """
                    INSERT INTO bet_settlements (
                        bet_id, settled_at, result, return_amount, profit_loss, metadata
                    ) VALUES (
                        :bet_id, :settled_at, :result, :return_amount, :profit_loss,
                        CAST(:metadata AS jsonb)
                    )
                    ON CONFLICT (bet_id)
                    DO UPDATE SET result = EXCLUDED.result,
                                  return_amount = EXCLUDED.return_amount,
                                  profit_loss = EXCLUDED.profit_loss
                    """
                ),
                {
                    "bet_id": bet_id,
                    "settled_at": placed_at,
                    "result": result,
                    "return_amount": return_amount,
                    "profit_loss": profit_loss,
                    "metadata": json.dumps(row, default=str),
                },
            )
        return bet_id

    def _import_movement(
        self, canonical: dict[str, str], row: dict[str, Any], account_id: int
    ) -> int | None:
        kind = self._value(row, canonical, "movement_type")
        amount = self._as_decimal(self._value(row, canonical, "stake"))
        if not kind or amount is None:
            return None
        normalized_kind = self._normalize(kind)
        aliases = {
            "deposito": "deposit",
            "deposit": "deposit",
            "saque": "withdrawal",
            "withdrawal": "withdrawal",
            "ajuste": "adjustment",
            "adjustment": "adjustment",
        }
        movement_type = aliases.get(normalized_kind)
        if movement_type is None:
            return None
        occurred_at = self._as_datetime(self._value(row, canonical, "date")) or datetime.now(UTC)
        signed_amount = -abs(amount) if movement_type == "withdrawal" else abs(amount)
        fingerprint = hashlib.sha256(
            f"movement|{movement_type}|{occurred_at.isoformat()}|{signed_amount}".encode()
        ).hexdigest()
        return self.db.execute(
            text(
                """
                INSERT INTO bankroll_movements (
                    account_id, movement_type, amount, balance_after, occurred_at,
                    source, source_fingerprint, created_at
                ) VALUES (
                    :account_id, :movement_type, :amount, :balance_after, :occurred_at,
                    'spreadsheet', :fingerprint, now()
                )
                ON CONFLICT (source_fingerprint)
                DO UPDATE SET balance_after = EXCLUDED.balance_after
                RETURNING id
                """
            ),
            {
                "account_id": account_id,
                "movement_type": movement_type,
                "amount": signed_amount,
                "balance_after": self._as_decimal(self._value(row, canonical, "balance")),
                "occurred_at": occurred_at,
                "fingerprint": fingerprint,
            },
        ).scalar_one()

    def _import_goal_timeline(
        self,
        match_id: int,
        timeline: str,
        home_id: int,
        away_id: int,
        home_name: str,
        away_name: str,
    ) -> None:
        self.db.execute(
            text("DELETE FROM match_events WHERE match_id = :match_id AND event_type = 'goal'"),
            {"match_id": match_id},
        )
        tokens = re.findall(r"([^,;|]+?)(\d{1,3})(?:\+(\d{1,2}))?['’]?", timeline)
        home_norm = self._normalize(home_name)
        away_norm = self._normalize(away_name)
        for label, minute, extra in tokens:
            label_norm = self._normalize(label)
            team_id = home_id if home_norm in label_norm or label_norm in home_norm else None
            if away_norm in label_norm or label_norm in away_norm:
                team_id = away_id
            self.db.execute(
                text(
                    """
                    INSERT INTO match_events (
                        match_id, team_id, period, minute, extra_minute,
                        event_type, metadata, created_at
                    ) VALUES (
                        :match_id, :team_id, NULL, :minute, :extra_minute,
                        'goal', '{}'::jsonb, now()
                    )
                    """
                ),
                {
                    "match_id": match_id,
                    "team_id": team_id,
                    "minute": int(minute),
                    "extra_minute": int(extra) if extra else None,
                },
            )

    def _ensure_account(self) -> int:
        return self.db.execute(
            text(
                """
                INSERT INTO betting_accounts (name, currency, initial_balance, created_at, updated_at)
                VALUES (:name, 'BRL', 100, now(), now())
                ON CONFLICT (name) DO UPDATE SET updated_at = now()
                RETURNING id
                """
            ),
            {"name": self.account_name},
        ).scalar_one()

    def _ensure_competition(self, name: str) -> int:
        existing = self.db.execute(
            text("SELECT id FROM competitions WHERE lower(name) = lower(:name) LIMIT 1"),
            {"name": name},
        ).scalar_one_or_none()
        if existing:
            return existing
        return self.db.execute(
            text(
                """
                INSERT INTO competitions (name, competition_type, created_at, updated_at)
                VALUES (:name, 'league', now(), now()) RETURNING id
                """
            ),
            {"name": name},
        ).scalar_one()

    def _ensure_season(self, competition_id: int, name: str) -> int:
        return self.db.execute(
            text(
                """
                INSERT INTO seasons (
                    competition_id, name, is_current, created_at, updated_at
                ) VALUES (:competition_id, :name, false, now(), now())
                ON CONFLICT (competition_id, name)
                DO UPDATE SET updated_at = now()
                RETURNING id
                """
            ),
            {"competition_id": competition_id, "name": name},
        ).scalar_one()

    def _ensure_team(self, name: str) -> int:
        existing = self.db.execute(
            text("SELECT id FROM teams WHERE lower(name) = lower(:name) LIMIT 1"),
            {"name": name.strip()},
        ).scalar_one_or_none()
        if existing:
            return existing
        return self.db.execute(
            text(
                """
                INSERT INTO teams (name, created_at, updated_at)
                VALUES (:name, now(), now()) RETURNING id
                """
            ),
            {"name": name.strip()},
        ).scalar_one()

    def _start_run(self, file_name: str, file_hash: str) -> int:
        return self.db.execute(
            text(
                """
                INSERT INTO import_runs (file_name, file_sha256, status)
                VALUES (:file_name, :file_hash, 'running') RETURNING id
                """
            ),
            {"file_name": file_name, "file_hash": file_hash},
        ).scalar_one()

    def _finish_run(self, result: SpreadsheetImportResult, status: str) -> None:
        self.db.execute(
            text(
                """
                UPDATE import_runs
                SET status = :status, finished_at = now(), rows_seen = :rows_seen,
                    rows_imported = :rows_imported, rows_skipped = :rows_skipped,
                    rows_failed = :rows_failed
                WHERE id = :run_id
                """
            ),
            {
                "status": status,
                "rows_seen": result.rows_seen,
                "rows_imported": result.rows_imported,
                "rows_skipped": result.rows_skipped,
                "rows_failed": result.rows_failed,
                "run_id": result.run_id,
            },
        )

    def _record_row(
        self,
        run_id: int,
        sheet_name: str,
        row_number: int,
        fingerprint: str,
        status: str,
        payload: dict[str, Any],
        entity_type: str | None = None,
        entity_id: int | None = None,
        error: str | None = None,
    ) -> None:
        self.db.execute(
            text(
                """
                INSERT INTO import_rows (
                    import_run_id, sheet_name, row_number, row_fingerprint, status,
                    entity_type, entity_id, error, raw_payload
                ) VALUES (
                    :run_id, :sheet_name, :row_number, :fingerprint, :status,
                    :entity_type, :entity_id, :error, CAST(:payload AS jsonb)
                )
                ON CONFLICT (row_fingerprint) DO NOTHING
                """
            ),
            {
                "run_id": run_id,
                "sheet_name": sheet_name,
                "row_number": row_number,
                "fingerprint": fingerprint,
                "status": status,
                "entity_type": entity_type,
                "entity_id": entity_id,
                "error": error,
                "payload": json.dumps(payload, default=str),
            },
        )

    def _fingerprint_exists(self, fingerprint: str) -> bool:
        return bool(
            self.db.execute(
                text("SELECT 1 FROM import_rows WHERE row_fingerprint = :fingerprint"),
                {"fingerprint": fingerprint},
            ).scalar_one_or_none()
        )

    @classmethod
    def _classify_sheet(cls, canonical: dict[str, str]) -> str:
        keys = set(canonical)
        if {"home", "away", "date"}.issubset(keys):
            return "match"
        if "stake" in keys and ("odd" in keys or "selection" in keys or "return" in keys):
            return "bet"
        if {"movement_type", "stake"}.issubset(keys):
            return "movement"
        return "unknown"

    @classmethod
    def _canonical_header_map(cls, headers: list[str]) -> dict[str, str]:
        mapping: dict[str, str] = {}
        for header in headers:
            for canonical, aliases in cls.HEADER_ALIASES.items():
                normalized_aliases = {cls._normalize(alias) for alias in aliases}
                if header in normalized_aliases:
                    mapping.setdefault(canonical, header)
        return mapping

    @staticmethod
    def _row_payload(headers: list[str], values: tuple[Any, ...]) -> dict[str, Any]:
        return {
            header: value
            for header, value in zip(headers, values, strict=False)
            if header and value is not None
        }

    @staticmethod
    def _fingerprint(sheet_name: str, payload: dict[str, Any]) -> str:
        canonical = json.dumps(payload, sort_keys=True, default=str, ensure_ascii=False)
        return hashlib.sha256(f"{sheet_name}|{canonical}".encode()).hexdigest()

    @staticmethod
    def _normalize(value: Any) -> str:
        if value is None:
            return ""
        text_value = unicodedata.normalize("NFKD", str(value).strip().lower())
        text_value = "".join(char for char in text_value if not unicodedata.combining(char))
        return re.sub(r"[^a-z0-9]+", "_", text_value).strip("_")

    @staticmethod
    def _value(row: dict[str, Any], canonical: dict[str, str], key: str) -> Any:
        header = canonical.get(key)
        return row.get(header) if header else None

    @staticmethod
    def _as_datetime(value: Any) -> datetime | None:
        if value is None:
            return None
        if isinstance(value, datetime):
            return value if value.tzinfo else value.replace(tzinfo=UTC)
        if isinstance(value, date):
            return datetime.combine(value, time.min, tzinfo=UTC)
        text_value = str(value).strip()
        for fmt in ("%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
            try:
                return datetime.strptime(text_value, fmt).replace(tzinfo=UTC)
            except ValueError:
                continue
        return None

    @staticmethod
    def _as_decimal(value: Any) -> Decimal | None:
        if value is None or value == "":
            return None
        if isinstance(value, Decimal):
            return value
        if isinstance(value, (int, float)):
            return Decimal(str(value))
        cleaned = re.sub(r"[^0-9,.-]", "", str(value)).replace(".", "").replace(",", ".")
        try:
            return Decimal(cleaned)
        except Exception:
            return None

    @staticmethod
    def _as_int(value: Any) -> int | None:
        try:
            return int(value) if value is not None else None
        except (TypeError, ValueError):
            return None

    @classmethod
    def _as_bool(cls, value: Any) -> bool:
        return cls._normalize(value) in {"1", "true", "sim", "yes", "live", "ao_vivo"}

    @classmethod
    def _normalize_bet_result(cls, value: Any) -> str | None:
        normalized = cls._normalize(value)
        if normalized in {"green", "ganhou", "ganha", "win", "won"}:
            return "Green"
        if normalized in {"red", "perdida", "perdeu", "lost", "loss"}:
            return "Red"
        return None

    @staticmethod
    def _parse_score(value: Any) -> tuple[int | None, int | None]:
        if value is None:
            return None, None
        match = re.search(r"(\d+)\s*[-xX:]\s*(\d+)", str(value))
        if not match:
            return None, None
        return int(match.group(1)), int(match.group(2))
